# -*- coding: utf-8 -*-
"""
Excel Export Service
تصدير البيانات إلى ملفات Excel (xlsx)

- تصدير الحوامل
- تصدير الزيارات
- تصدير الوصفات
- تصدير التحصينات
- تصدير الكل (ملف واحد، عدة sheets)
"""
import io
from datetime import datetime, date
from typing import List, Dict, Optional, Any

from core.db import db, select


# ============ HELPERS ============

def _safe_str(v: Any) -> str:
    """Convert any value to safe string for Excel"""
    if v is None:
        return ""
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, bool):
        return "نعم" if v else "لا"
    if isinstance(v, (dict, list)):
        import json
        try:
            return json.dumps(v, ensure_ascii=False)
        except Exception:
            return str(v)
    return str(v)


def _safe_date(v: Any) -> str:
    """Format datetime/date for Excel"""
    if v is None:
        return ""
    if isinstance(v, str):
        # Try to parse ISO
        try:
            if "T" in v:
                return datetime.fromisoformat(v.replace("Z", "+00:00")).strftime("%Y-%m-%d")
            return v[:10]
        except Exception:
            return v
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    return str(v)


def _fetch_all(table: str, filters: Dict, order_by: str = "created_at", desc: bool = True) -> List[Dict]:
    """
    Fetch ALL rows from a Supabase table (bypasses 1000-row limit).
    Uses pagination with range().
    """
    PAGE_SIZE = 1000
    all_rows = []
    offset = 0
    while True:
        try:
            q = db().table(table).select("*")
            for k, v in filters.items():
                q = q.eq(k, v)
            if order_by:
                q = q.order(order_by, desc=desc)
            q = q.range(offset, offset + PAGE_SIZE - 1)
            rows = q.execute().data or []
            all_rows.extend(rows)
            if len(rows) < PAGE_SIZE:
                break
            offset += PAGE_SIZE
        except Exception:
            break
    return all_rows


def _get_patients_map(clinic_id: str) -> Dict[str, Dict]:
    """Get {patient_id: patient_row} for name lookups"""
    rows = _fetch_all("patients", {"clinic_id": clinic_id}, order_by=None)
    return {p["id"]: p for p in rows}


def _get_users_map(clinic_id: str) -> Dict[str, Dict]:
    """Get {user_id: user_row} for doctor name lookups"""
    rows = _fetch_all("users", {"clinic_id": clinic_id}, order_by=None)
    return {u["id"]: u for u in rows}


def _calc_gestational_weeks(lmp: Any, ref_date: Any = None) -> str:
    """Calculate gestational age from LMP"""
    if not lmp:
        return ""
    try:
        if isinstance(lmp, str):
            lmp_date = datetime.fromisoformat(lmp[:10]).date()
        elif isinstance(lmp, date):
            lmp_date = lmp
        else:
            return ""

        if ref_date:
            if isinstance(ref_date, str):
                ref = datetime.fromisoformat(ref_date[:10]).date()
            elif isinstance(ref_date, date):
                ref = ref_date
            else:
                ref = date.today()
        else:
            ref = date.today()

        days = (ref - lmp_date).days
        if days < 0:
            return ""
        weeks = days // 7
        return "{} أسبوع".format(weeks)
    except Exception:
        return ""


# ============ EXPORT FUNCTIONS ============

def export_patients(clinic_id: str, include_inactive: bool = False) -> bytes:
    """Export patients to Excel"""
    import pandas as pd

    filters = {"clinic_id": clinic_id}
    rows = _fetch_all("patients", filters, order_by="created_at", desc=True)

    if not include_inactive:
        rows = [r for r in rows if r.get("is_active", True)]

    data = []
    for p in rows:
        data.append({
            "الاسم الكامل": _safe_str(p.get("full_name")),
            "الهاتف": _safe_str(p.get("phone")),
            "الرقم الوطني": _safe_str(p.get("national_id")),
            "تاريخ الميلاد": _safe_date(p.get("date_of_birth")),
            "فصيلة الدم": _safe_str(p.get("blood_type")),
            "LMP": _safe_date(p.get("lmp")),
            "EDD": _safe_date(p.get("edd")),
            "أسابيع الحمل": _calc_gestational_weeks(p.get("lmp")),
            "Gravida": p.get("gravida") if p.get("gravida") is not None else "",
            "Para": p.get("para") if p.get("para") is not None else "",
            "مستوى الخطر": _safe_str(p.get("risk_level")),
            "ملاحظات": _safe_str(p.get("notes")),
            "نشط": "نعم" if p.get("is_active") else "لا",
            "تاريخ التسجيل": _safe_date(p.get("created_at")),
        })

    df = pd.DataFrame(data) if data else pd.DataFrame(columns=[
        "الاسم الكامل", "الهاتف", "الرقم الوطني", "تاريخ الميلاد", "فصيلة الدم",
        "LMP", "EDD", "أسابيع الحمل", "Gravida", "Para", "مستوى الخطر",
        "ملاحظات", "نشط", "تاريخ التسجيل"
    ])

    return _df_to_xlsx(df, "الحوامل")


def export_visits(clinic_id: str, date_from=None, date_to=None) -> bytes:
    """Export visits to Excel"""
    import pandas as pd

    rows = _fetch_all("visits", {"clinic_id": clinic_id}, order_by="visit_date", desc=True)

    # Date filter
    if date_from or date_to:
        filtered = []
        for v in rows:
            vd = v.get("visit_date", "")
            if vd:
                vd_str = _safe_date(vd)
                if date_from and vd_str < str(date_from):
                    continue
                if date_to and vd_str > str(date_to):
                    continue
            filtered.append(v)
        rows = filtered

    patients_map = _get_patients_map(clinic_id)
    users_map = _get_users_map(clinic_id)

    data = []
    for v in rows:
        pid = v.get("patient_id")
        did = v.get("doctor_id")
        patient = patients_map.get(pid, {})
        doctor = users_map.get(did, {})

        data.append({
            "التاريخ": _safe_date(v.get("visit_date")),
            "الحامل": _safe_str(patient.get("full_name")),
            "هاتف الحامل": _safe_str(patient.get("phone")),
            "الطبيبة": _safe_str(doctor.get("full_name")),
            "أسابيع الحمل": v.get("gestational_age_weeks") if v.get("gestational_age_weeks") is not None else "",
            "الوزن (كغ)": v.get("weight_kg") if v.get("weight_kg") is not None else "",
            "ضغط انقباضي": v.get("bp_systolic") if v.get("bp_systolic") is not None else "",
            "ضغط انبساطي": v.get("bp_diastolic") if v.get("bp_diastolic") is not None else "",
            "ارتفاع الرحم (سم)": v.get("fundal_height_cm") if v.get("fundal_height_cm") is not None else "",
            "نبض الجنين": v.get("fetal_heart_rate") if v.get("fetal_heart_rate") is not None else "",
            "ملاحظات": _safe_str(v.get("notes")),
            "تحليل AI": _safe_str(v.get("ai_risk_analysis"))[:500],
        })

    df = pd.DataFrame(data) if data else pd.DataFrame(columns=[
        "التاريخ", "الحامل", "هاتف الحامل", "الطبيبة", "أسابيع الحمل",
        "الوزن (كغ)", "ضغط انقباضي", "ضغط انبساطي", "ارتفاع الرحم (سم)",
        "نبض الجنين", "ملاحظات", "تحليل AI"
    ])

    return _df_to_xlsx(df, "الزيارات")


def export_prescriptions(clinic_id: str, date_from=None, date_to=None) -> bytes:
    """Export prescriptions to Excel"""
    import pandas as pd

    rows = _fetch_all("prescriptions", {"clinic_id": clinic_id}, order_by="issued_at", desc=True)

    # Date filter
    if date_from or date_to:
        filtered = []
        for p in rows:
            pd_str = _safe_date(p.get("issued_at"))
            if date_from and pd_str < str(date_from):
                continue
            if date_to and pd_str > str(date_to):
                continue
            filtered.append(p)
        rows = filtered

    patients_map = _get_patients_map(clinic_id)
    users_map = _get_users_map(clinic_id)

    data = []
    for p in rows:
        patient = patients_map.get(p.get("patient_id"), {})
        doctor = users_map.get(p.get("doctor_id"), {})

        # Medications (jsonb)
        meds = p.get("medications")
        meds_str = ""
        if isinstance(meds, list):
            try:
                parts = []
                for m in meds:
                    if isinstance(m, dict):
                        name = m.get("name") or m.get("drug") or ""
                        dose = m.get("dose") or m.get("dosage") or ""
                        freq = m.get("frequency") or ""
                        dur = m.get("duration") or ""
                        parts.append("{} {} {} {}".format(name, dose, freq, dur).strip())
                    else:
                        parts.append(str(m))
                meds_str = " | ".join(parts)
            except Exception:
                meds_str = str(meds)
        elif meds:
            meds_str = str(meds)

        data.append({
            "التاريخ": _safe_date(p.get("issued_at")),
            "الحامل": _safe_str(patient.get("full_name")),
            "هاتف الحامل": _safe_str(patient.get("phone")),
            "الطبيبة": _safe_str(doctor.get("full_name")),
            "التشخيص": _safe_str(p.get("diagnosis")),
            "الأدوية": meds_str,
            "التعليمات": _safe_str(p.get("instructions")),
            "المتابعة (يوم)": p.get("follow_up_days") if p.get("follow_up_days") is not None else "",
            "ملاحظات": _safe_str(p.get("notes")),
        })

    df = pd.DataFrame(data) if data else pd.DataFrame(columns=[
        "التاريخ", "الحامل", "هاتف الحامل", "الطبيبة", "التشخيص",
        "الأدوية", "التعليمات", "المتابعة (يوم)", "ملاحظات"
    ])

    return _df_to_xlsx(df, "الوصفات")


def export_vaccinations(clinic_id: str, date_from=None, date_to=None) -> bytes:
    """Export vaccinations to Excel"""
    import pandas as pd

    rows = _fetch_all("vaccinations", {"clinic_id": clinic_id}, order_by="created_at", desc=True)

    # Date filter (on given_at)
    if date_from or date_to:
        filtered = []
        for v in rows:
            gd = _safe_date(v.get("given_at")) or _safe_date(v.get("created_at"))
            if date_from and gd < str(date_from):
                continue
            if date_to and gd > str(date_to):
                continue
            filtered.append(v)
        rows = filtered

    patients_map = _get_patients_map(clinic_id)
    users_map = _get_users_map(clinic_id)

    data = []
    for v in rows:
        patient = patients_map.get(v.get("patient_id"), {})
        doctor = users_map.get(v.get("doctor_id"), {})

        data.append({
            "الحامل": _safe_str(patient.get("full_name")),
            "هاتف الحامل": _safe_str(patient.get("phone")),
            "الطبيبة": _safe_str(doctor.get("full_name")),
            "اسم اللقاح": _safe_str(v.get("vaccine_name")),
            "رقم الجرعة": v.get("dose_number") if v.get("dose_number") is not None else "",
            "تاريخ الإعطاء": _safe_date(v.get("given_at")),
            "الموعد المجدول": _safe_date(v.get("scheduled_for")),
            "الحالة": _safe_str(v.get("status")),
            "رقم الدفعة": _safe_str(v.get("batch_number")),
            "الموقع": _safe_str(v.get("site")),
            "ملاحظات": _safe_str(v.get("notes")),
            "تاريخ التسجيل": _safe_date(v.get("created_at")),
        })

    df = pd.DataFrame(data) if data else pd.DataFrame(columns=[
        "الحامل", "هاتف الحامل", "الطبيبة", "اسم اللقاح", "رقم الجرعة",
        "تاريخ الإعطاء", "الموعد المجدول", "الحالة", "رقم الدفعة", "الموقع",
        "ملاحظات", "تاريخ التسجيل"
    ])

    return _df_to_xlsx(df, "التحصينات")


def export_all(clinic_id: str, date_from=None, date_to=None) -> bytes:
    """
    Export everything to a single Excel file with multiple sheets.
    Returns: bytes of xlsx file
    """
    import pandas as pd

    patients_map = _get_patients_map(clinic_id)
    users_map = _get_users_map(clinic_id)

    # ---- Patients sheet ----
    pat_rows = _fetch_all("patients", {"clinic_id": clinic_id}, order_by="created_at", desc=True)
    pat_data = [{
        "الاسم الكامل": _safe_str(p.get("full_name")),
        "الهاتف": _safe_str(p.get("phone")),
        "الرقم الوطني": _safe_str(p.get("national_id")),
        "تاريخ الميلاد": _safe_date(p.get("date_of_birth")),
        "فصيلة الدم": _safe_str(p.get("blood_type")),
        "LMP": _safe_date(p.get("lmp")),
        "EDD": _safe_date(p.get("edd")),
        "أسابيع الحمل": _calc_gestational_weeks(p.get("lmp")),
        "Gravida": p.get("gravida") if p.get("gravida") is not None else "",
        "Para": p.get("para") if p.get("para") is not None else "",
        "مستوى الخطر": _safe_str(p.get("risk_level")),
        "ملاحظات": _safe_str(p.get("notes")),
        "نشط": "نعم" if p.get("is_active") else "لا",
        "تاريخ التسجيل": _safe_date(p.get("created_at")),
    } for p in pat_rows]
    df_pat = pd.DataFrame(pat_data)

    # ---- Visits sheet ----
    vis_rows = _fetch_all("visits", {"clinic_id": clinic_id}, order_by="visit_date", desc=True)
    vis_data = []
    for v in vis_rows:
        patient = patients_map.get(v.get("patient_id"), {})
        doctor = users_map.get(v.get("doctor_id"), {})
        vis_data.append({
            "التاريخ": _safe_date(v.get("visit_date")),
            "الحامل": _safe_str(patient.get("full_name")),
            "هاتف الحامل": _safe_str(patient.get("phone")),
            "الطبيبة": _safe_str(doctor.get("full_name")),
            "أسابيع الحمل": v.get("gestational_age_weeks") if v.get("gestational_age_weeks") is not None else "",
            "الوزن (كغ)": v.get("weight_kg") if v.get("weight_kg") is not None else "",
            "ضغط انقباضي": v.get("bp_systolic") if v.get("bp_systolic") is not None else "",
            "ضغط انبساطي": v.get("bp_diastolic") if v.get("bp_diastolic") is not None else "",
            "ارتفاع الرحم (سم)": v.get("fundal_height_cm") if v.get("fundal_height_cm") is not None else "",
            "نبض الجنين": v.get("fetal_heart_rate") if v.get("fetal_heart_rate") is not None else "",
            "ملاحظات": _safe_str(v.get("notes")),
        })
    df_vis = pd.DataFrame(vis_data)

    # ---- Prescriptions sheet ----
    pres_rows = _fetch_all("prescriptions", {"clinic_id": clinic_id}, order_by="issued_at", desc=True)
    pres_data = []
    for p in pres_rows:
        patient = patients_map.get(p.get("patient_id"), {})
        doctor = users_map.get(p.get("doctor_id"), {})
        meds = p.get("medications")
        meds_str = ""
        if isinstance(meds, list):
            try:
                parts = []
                for m in meds:
                    if isinstance(m, dict):
                        name = m.get("name") or m.get("drug") or ""
                        dose = m.get("dose") or m.get("dosage") or ""
                        freq = m.get("frequency") or ""
                        dur = m.get("duration") or ""
                        parts.append("{} {} {} {}".format(name, dose, freq, dur).strip())
                    else:
                        parts.append(str(m))
                meds_str = " | ".join(parts)
            except Exception:
                meds_str = str(meds)
        elif meds:
            meds_str = str(meds)

        pres_data.append({
            "التاريخ": _safe_date(p.get("issued_at")),
            "الحامل": _safe_str(patient.get("full_name")),
            "هاتف الحامل": _safe_str(patient.get("phone")),
            "الطبيبة": _safe_str(doctor.get("full_name")),
            "التشخيص": _safe_str(p.get("diagnosis")),
            "الأدوية": meds_str,
            "التعليمات": _safe_str(p.get("instructions")),
            "المتابعة (يوم)": p.get("follow_up_days") if p.get("follow_up_days") is not None else "",
            "ملاحظات": _safe_str(p.get("notes")),
        })
    df_pres = pd.DataFrame(pres_data)

    # ---- Vaccinations sheet ----
    vacc_rows = _fetch_all("vaccinations", {"clinic_id": clinic_id}, order_by="created_at", desc=True)
    vacc_data = []
    for v in vacc_rows:
        patient = patients_map.get(v.get("patient_id"), {})
        doctor = users_map.get(v.get("doctor_id"), {})
        vacc_data.append({
            "الحامل": _safe_str(patient.get("full_name")),
            "هاتف الحامل": _safe_str(patient.get("phone")),
            "الطبيبة": _safe_str(doctor.get("full_name")),
            "اسم اللقاح": _safe_str(v.get("vaccine_name")),
            "رقم الجرعة": v.get("dose_number") if v.get("dose_number") is not None else "",
            "تاريخ الإعطاء": _safe_date(v.get("given_at")),
            "الموعد المجدول": _safe_date(v.get("scheduled_for")),
            "الحالة": _safe_str(v.get("status")),
            "رقم الدفعة": _safe_str(v.get("batch_number")),
            "الموقع": _safe_str(v.get("site")),
            "ملاحظات": _safe_str(v.get("notes")),
        })
    df_vacc = pd.DataFrame(vacc_data)

    # ---- Build xlsx with multiple sheets ----
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        _write_sheet(writer, df_pat, "الحوامل", [
            "الاسم الكامل", "الهاتف", "الرقم الوطني", "تاريخ الميلاد", "فصيلة الدم",
            "LMP", "EDD", "أسابيع الحمل", "Gravida", "Para", "مستوى الخطر",
            "ملاحظات", "نشط", "تاريخ التسجيل"
        ])
        _write_sheet(writer, df_vis, "الزيارات", [
            "التاريخ", "الحامل", "هاتف الحامل", "الطبيبة", "أسابيع الحمل",
            "الوزن (كغ)", "ضغط انقباضي", "ضغط انبساطي", "ارتفاع الرحم (سم)",
            "نبض الجنين", "ملاحظات"
        ])
        _write_sheet(writer, df_pres, "الوصفات", [
            "التاريخ", "الحامل", "هاتف الحامل", "الطبيبة", "التشخيص",
            "الأدوية", "التعليمات", "المتابعة (يوم)", "ملاحظات"
        ])
        _write_sheet(writer, df_vacc, "التحصينات", [
            "الحامل", "هاتف الحامل", "الطبيبة", "اسم اللقاح", "رقم الجرعة",
            "تاريخ الإعطاء", "الموعد المجدول", "الحالة", "رقم الدفعة",
            "الموقع", "ملاحظات"
        ])

    output.seek(0)
    return output.getvalue()


# ============ INTERNAL ============

def _df_to_xlsx(df, sheet_name: str) -> bytes:
    """Convert DataFrame to xlsx bytes with formatting"""
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name=sheet_name)
        _format_worksheet(writer.sheets[sheet_name], df)
    output.seek(0)
    return output.getvalue()


def _write_sheet(writer, df, sheet_name: str, columns: list):
    """Write a sheet with fixed columns (even if empty)"""
    import pandas as pd
    if df.empty:
        df = pd.DataFrame(columns=columns)
    else:
        # Ensure all columns exist
        for c in columns:
            if c not in df.columns:
                df[c] = ""
        df = df[columns]
    df.to_excel(writer, index=False, sheet_name=sheet_name)
    _format_worksheet(writer.sheets[sheet_name], df)


def _format_worksheet(ws, df):
    """Apply formatting: header style, widths, freeze, RTL"""
    try:
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

        # Header style
        header_font = Font(bold=True, color="FFFFFF", size=11)
        header_fill = PatternFill(start_color="16A34A", end_color="16A34A", fill_type="solid")
        header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

        # Apply to header row
        for col_idx, _ in enumerate(df.columns, start=1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align

        # Column widths (approximate)
        for col_idx, col_name in enumerate(df.columns, start=1):
            max_len = len(str(col_name))
            for row_idx in range(2, min(ws.max_row + 1, 200)):
                v = ws.cell(row=row_idx, column=col_idx).value
                if v is not None:
                    max_len = max(max_len, min(len(str(v)), 60))
            width = min(max(max_len + 3, 12), 50)
            ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = width

        # Freeze header row
        ws.freeze_panes = "A2"

        # RTL (right-to-left) for Arabic
        ws.sheet_view.rightToLeft = True

        # Row height for header
        ws.row_dimensions[1].height = 28
    except Exception:
        pass


# ============ METADATA ============

EXPORTS = {
    "patients": {
        "label": "👩 الحوامل",
        "description": "قائمة كاملة بالحوامل المسجلات",
        "func": export_patients,
        "has_date_filter": False,
    },
    "visits": {
        "label": "🩺 الزيارات",
        "description": "جميع الزيارات الطبية المسجلة",
        "func": export_visits,
        "has_date_filter": True,
    },
    "prescriptions": {
        "label": "💊 الوصفات",
        "description": "جميع الوصفات الطبية",
        "func": export_prescriptions,
        "has_date_filter": True,
    },
    "vaccinations": {
        "label": "💉 التحصينات",
        "description": "جميع التحصينات",
        "func": export_vaccinations,
        "has_date_filter": True,
    },
}
